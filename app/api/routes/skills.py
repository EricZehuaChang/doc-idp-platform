"""Skills API (design v0.2 §7 group 2): CRUD + versioning + studio services
(probe pre-annotation / dry-run side-by-side / YAML import-export / golden
check). All studio interactions are API-first — the M2 UI sits on these.
"""
import asyncio
import hashlib
from pathlib import Path

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import PlainTextResponse
from pydantic import BaseModel
from sqlalchemy import func, select

from app.billing import engine as billing
from app.config import get_settings
from app.extraction.provider_client import ProviderError
from app.db import session_factory
from app.models import GoldenSample, Skill, SkillVersion, Transaction
from app.parsers.base import UDR
from app.parsers.router import parse_document
from app.skillengine import catalog, studio
from app.skillengine.schema import SkillPackage
from app.tenancy import current_actor, current_tenant


async def _enforce_skill_seat(s, tenant: str) -> None:
    """Plan skill cap (§12.3); Owner Root is exempt from feature gates (§12.7)."""
    if current_actor().get("unlimited"):
        return
    try:
        await billing.enforce_skill_cap(s, tenant)
    except billing.EntitlementExceeded as e:
        raise HTTPException(403, f"当前套餐({e.plan})技能数已达上限 {e.limit},"
                                 "请删除闲置技能或联系平台升级套餐")

router = APIRouter(prefix="/api/v1/skills", tags=["skills"])


async def _parse_upload(up: UploadFile, tenant: str) -> UDR:
    """Persist an uploaded sample then parse it to UDR (thread offload)."""
    _path, udr = await _parse_upload_kept(up, tenant)
    return udr


async def _parse_upload_kept(up: UploadFile, tenant: str,
                             pinned: str | None = None) -> tuple[str, UDR]:
    """Same, but hands back the stored path — the vision dry-run rasterizes the
    original, which the UDR alone cannot reconstruct.

    `pinned` must be the skill's own parser (2026-08-27): the studio used to
    auto-route regardless, so a skill pinned to a non-default parser was tried
    and gated against text a different engine produced. Dry-run is exactly where
    a new parser gets evaluated — silently substituting one makes the evaluation
    meaningless — and golden-check is the publish gate, where it makes the gate
    lie about what production will do."""
    blob = await up.read()
    digest = hashlib.sha256(blob).hexdigest()[:16]
    suffix = Path(up.filename or "sample").suffix.lower()
    store = Path(get_settings().data_dir) / "samples" / tenant
    store.mkdir(parents=True, exist_ok=True)
    path = store / f"{digest}{suffix}"
    path.write_bytes(blob)
    return str(path), await asyncio.to_thread(parse_document, str(path), pinned)


class SkillCreate(BaseModel):
    package: SkillPackage
    changelog: str = ""


@router.post("", status_code=201)
async def create_skill(payload: SkillCreate):
    tenant = current_tenant()
    pkg = payload.package
    sf = session_factory()
    async with sf() as s:
        if await s.get(Skill, pkg.skill_code) is not None:
            raise HTTPException(409, f"skill exists: {pkg.skill_code}")
        await _enforce_skill_seat(s, tenant)
        s.add(Skill(code=pkg.skill_code, tenant_id=tenant,
                    name=pkg.name or pkg.skill_code, kind=pkg.kind))
        s.add(SkillVersion(tenant_id=tenant, skill_code=pkg.skill_code, version=1,
                           status="draft", package=pkg.model_dump(),
                           changelog=payload.changelog or "initial draft"))
        await s.commit()
    return {"skill_code": pkg.skill_code, "version": 1, "status": "draft"}


@router.get("")
async def list_skills(state: str | None = None):
    """Skill roster. `state` selects a lifecycle bucket: the default view hides
    deleted skills, `state=deleted` lists the archive (soft-deleted skills keep
    every version and can be restored — 需求7 skill archive)."""
    tenant = current_tenant()
    if state and state not in ("active", "disabled", "deleted"):
        raise HTTPException(400, "state 只能是 active / disabled / deleted")
    sf = session_factory()
    async with sf() as s:
        cond = [Skill.tenant_id == tenant]
        if state:
            cond.append(Skill.state == state)
        else:
            cond.append(Skill.state != "deleted")
        rows = (await s.execute(select(Skill).where(*cond))).scalars().all()
        # highest published version per skill: /process rejects a skill without
        # one (400), so the upload page must be able to hide those up front
        pub = dict((await s.execute(
            select(SkillVersion.skill_code, func.max(SkillVersion.version))
            .where(SkillVersion.skill_code.in_([r.code for r in rows]),
                   SkillVersion.status == "published")
            .group_by(SkillVersion.skill_code))).all())
        return [{"skill_code": r.code, "name": r.name, "kind": r.kind, "state": r.state,
                 "published_version": pub.get(r.code)}
                for r in rows]


@router.get("/model-options")
async def model_options():
    """Provider + parser catalogue for the skill editor's pickers.

    Deliberately not the BYOK endpoint (`/settings/providers`): that one is
    admin-only and reports key presence. Editors are used by operators, and the
    only thing the model/parser fields need is the list of configured names —
    so this returns names and model ids and nothing else. Declared before
    GET /{skill_code} so "model-options" is not read as a skill code.
    """
    from app.config import load_parsers, load_providers
    from app.extraction import custom_providers

    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        stored = await custom_providers.load_raw(s, tenant)
    cfg = load_providers()
    out = [{"name": name, "model": p.model, "active": name == cfg["active"],
            "custom": False, "vision": False}
           for name, p in cfg["providers"].items()]
    # console-registered channels are pickable exactly like the built-in ones —
    # that is the whole point of registering them
    out += [{"name": n, "model": e.get("model") or "", "active": False,
             "custom": True, "vision": bool(e.get("vision"))}
            for n, e in sorted(stored.items()) if isinstance(e, dict)]
    return {
        "providers": out,
        "fallback_chain": list(cfg.get("fallback") or []),
        # parsers carry a description so the editor can show what each vlm-*
        # routing name actually resolves to (model + capability), not just the key
        "parsers": [{"name": p.name, "type": p.type, "description": p.description}
                    for p in load_parsers()["parsers"].values()],
    }


@router.get("/templates")
async def list_skill_templates():
    """Built-in starter templates for the "new skill" gallery (onboarding P0).

    Read-only and tenant-independent — templates ship with the build, so there
    is nothing here to scope. Declared before GET /{skill_code} so "templates"
    is not read as a skill code.
    """
    return {"templates": catalog.list_templates()}


@router.post("/templates/{template_id}/import", status_code=201)
async def import_skill_template(template_id: str):
    """Instantiate a built-in template as a draft skill for this tenant.

    Same landing state as POST /import (draft v1, editor opens on it) — the
    only difference is where the package came from. The skill code is derived
    per tenant because `Skill.code` is a global primary key: handing out the
    template id verbatim would 409 for every tenant after the first.
    """
    pkg = catalog.get_template(template_id)
    if pkg is None:
        raise HTTPException(404, f"模板不存在: {template_id}")
    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        await _enforce_skill_seat(s, tenant)   # a template still costs a seat
        code = None
        for cand in catalog.candidate_codes(template_id, tenant):
            if await s.get(Skill, cand) is None:
                code = cand
                break
        if code is None:
            raise HTTPException(409, f"该模板已导入过多次，请先清理旧的「{pkg.name}」技能")
        pkg.skill_code = code
        s.add(Skill(code=code, tenant_id=tenant, name=pkg.name or code, kind=pkg.kind))
        s.add(SkillVersion(tenant_id=tenant, skill_code=code, version=1,
                           status="draft", package=pkg.model_dump(),
                           changelog=f"从内置模板「{pkg.name}」创建"))
        await s.commit()
    return {"skill_code": code, "version": 1, "status": "draft",
            "template_id": template_id, "name": pkg.name}


@router.get("/{skill_code}")
async def get_skill(skill_code: str, version: int | None = None):
    """Skill detail; ?version=N loads that version's package into the editor
    (the version rail switches between drafts/published/archived)."""
    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        skill = await s.get(Skill, skill_code)
        if skill is None or skill.tenant_id != tenant:
            raise HTTPException(404, "skill not found")
        versions = (await s.execute(
            select(SkillVersion).where(SkillVersion.skill_code == skill_code)
            .order_by(SkillVersion.version))).scalars().all()
        selected = None
        if version is not None:
            selected = next((v for v in versions if v.version == version), None)
            if selected is None:
                raise HTTPException(404, "version not found")
        elif versions:
            selected = versions[-1]
        return {"skill_code": skill.code, "name": skill.name, "kind": skill.kind,
                "state": skill.state,
                "versions": [{"version": v.version, "status": v.status,
                              "changelog": v.changelog,
                              "created_at": v.created_at.isoformat()} for v in versions],
                "selected_version": selected.version if selected else None,
                "latest_package": selected.package if selected else None}


@router.delete("/{skill_code}")
async def delete_skill(skill_code: str):
    """Soft delete (state machine, §5.1): running tasks keep their pinned
    versions; the skill just disappears from lists and submission."""
    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        skill = await s.get(Skill, skill_code)
        if skill is None or skill.tenant_id != tenant:
            raise HTTPException(404, "skill not found")
        skill.state = "deleted"
        await s.commit()
    return {"skill_code": skill_code, "state": "deleted"}


@router.post("/{skill_code}/restore")
async def restore_skill(skill_code: str):
    """Bring an archived (soft-deleted) skill back. Nothing was dropped by the
    delete — versions, golden samples and history all stayed — so restoring is
    a state flip back to active (需求7: 删除的技能可找回)."""
    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        skill = await s.get(Skill, skill_code)
        if skill is None or skill.tenant_id != tenant:
            raise HTTPException(404, "skill not found")
        if skill.state != "deleted":
            raise HTTPException(409, f"技能当前状态为 {skill.state}，无需恢复")
        await _enforce_skill_seat(s, tenant)   # a restored skill is a seat again
        skill.state = "active"
        await s.commit()
    return {"skill_code": skill_code, "state": "active"}


class DraftUpdate(BaseModel):
    package: SkillPackage
    changelog: str = ""


@router.post("/{skill_code}/versions", status_code=201)
async def new_draft(skill_code: str, payload: DraftUpdate):
    """Published versions are immutable (§5.1) — any change opens a new draft."""
    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        skill = await s.get(Skill, skill_code)
        if skill is None or skill.tenant_id != tenant:
            raise HTTPException(404, "skill not found")
        latest = (await s.execute(
            select(SkillVersion).where(SkillVersion.skill_code == skill_code)
            .order_by(SkillVersion.version.desc()))).scalars().first()
        next_ver = (latest.version + 1) if latest else 1
        s.add(SkillVersion(tenant_id=tenant, skill_code=skill_code, version=next_ver,
                           status="draft", package=payload.package.model_dump(),
                           changelog=payload.changelog))
        await s.commit()
    return {"skill_code": skill_code, "version": next_ver, "status": "draft"}


@router.put("/{skill_code}/versions/{version}")
async def update_draft(skill_code: str, version: int, payload: DraftUpdate):
    """Save into an existing draft, in place.

    Business rule (P04, 2026-08-26): 「保存」 updates the draft the editor is on;
    only 「新建版本」 (POST .../versions) and 「发布」 move the version pointer.
    Before this endpoint existed the editor's save button called the create-draft
    route, so every save minted vN+1 and polluted the version list.
    Published/archived versions stay immutable (§5.1) — saving onto one is a 409
    and the caller is expected to branch a new draft instead.
    """
    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        skill = await s.get(Skill, skill_code)
        if skill is None or skill.tenant_id != tenant:
            raise HTTPException(404, "skill not found")
        row = (await s.execute(
            select(SkillVersion).where(SkillVersion.skill_code == skill_code,
                                       SkillVersion.tenant_id == tenant,
                                       SkillVersion.version == version))).scalar_one_or_none()
        if row is None:
            raise HTTPException(404, "version not found")
        if row.status != "draft":
            raise HTTPException(409, f"v{version} 已{row.status}，不可原地修改；请另存为新草稿")
        row.package = payload.package.model_dump()
        # empty changelog means "not edited in this save" — keep the stored one
        # rather than blanking the version note (P05)
        if payload.changelog:
            row.changelog = payload.changelog
        skill.name = payload.package.name or skill.name
        await s.commit()
    return {"skill_code": skill_code, "version": version, "status": "draft"}


# a transaction whose pipeline has not finished still reloads its package by
# (skill_code, version) — anything past these two is history the runner never
# reads again (2026-09-04)
_LIVE_TXN = ("queued", "processing")


@router.delete("/{skill_code}/versions/{version}")
async def delete_version(skill_code: str, version: int):
    """Delete the version on screen (需求8: 「删除当前版本」与「删除技能」是两个动作).

    Widened 2026-09-04: draft-only was too narrow to be the button people
    actually wanted — every archived version in a real console is one a task
    once ran, so nothing in the version rail could ever be cleaned up. What
    stays protected is only what is still load-bearing:

    - the published version, because /process resolves submissions to it;
    - the last remaining version, because a skill with none can neither be
      opened in the editor nor submitted against;
    - any version an unfinished task pinned, because runner._load_package
      reads the row back by version number while the pipeline is running.

    A completed task keeps its results and its recorded version number either
    way: nothing outside /process and the runner reads a version row.
    """
    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        skill = await s.get(Skill, skill_code)
        if skill is None or skill.tenant_id != tenant:
            raise HTTPException(404, "skill not found")
        row = (await s.execute(
            select(SkillVersion).where(SkillVersion.skill_code == skill_code,
                                       SkillVersion.tenant_id == tenant,
                                       SkillVersion.version == version))).scalar_one_or_none()
        if row is None:
            raise HTTPException(404, "version not found")
        if row.status == "published":
            raise HTTPException(409, f"v{version} 是当前发布版本，不可删除；"
                                     "请先发布其它版本，或使用「删除技能」")
        remaining = (await s.execute(
            select(func.count(SkillVersion.id))
            .where(SkillVersion.skill_code == skill_code,
                   SkillVersion.tenant_id == tenant))).scalar_one()
        if remaining <= 1:
            raise HTTPException(409, "每个技能至少保留一个版本；要移除整个技能请使用「删除技能」")
        live = (await s.execute(
            select(func.count(Transaction.id))
            .where(Transaction.tenant_id == tenant,
                   Transaction.skill_code == skill_code,
                   Transaction.skill_version == version,
                   Transaction.status.in_(_LIVE_TXN)))).scalar_one()
        if live:
            raise HTTPException(409, f"v{version} 还有 {live} 个任务在跑，等跑完再删")
        await s.delete(row)
        await s.commit()
    return {"skill_code": skill_code, "version": version, "status": "deleted"}


@router.post("/probe")
async def probe(file: UploadFile = File(...), provider: str | None = Form(default=None)):
    """Sample -> LLM-drafted field list (§5.1 step 2: pre-annotation).
    The editor opens 80% filled instead of blank."""
    udr = await _parse_upload(file, current_tenant())
    await _warm_byok()
    try:
        out = await asyncio.to_thread(studio.probe, udr, provider)
    except ProviderError as e:
        raise HTTPException(502, f"模型通道暂时不可用，请稍后重试：{str(e)[:200]}")
    draft = out["draft"]
    fields = studio.draft_to_fields(draft if isinstance(draft, dict) else {})
    return {"doc_type": (draft or {}).get("doc_type", ""),
            "fields": [f.model_dump() for f in fields],
            "raw_draft": draft, "provider_used": out["provider_used"]}


class TextDraftBody(BaseModel):
    text: str
    provider: str | None = None


@router.post("/draft-from-text")
async def draft_from_text(body: TextDraftBody):
    """Rule import channel 2 (§5.1): natural-language requirement -> drafted
    fields, prefilled into the editor for human review before saving."""
    if not body.text.strip():
        raise HTTPException(400, "描述不能为空")
    await _warm_byok()
    try:
        out = await asyncio.to_thread(studio.draft_from_text, body.text, body.provider)
    except ProviderError as e:
        raise HTTPException(502, f"模型通道暂时不可用，请稍后重试：{str(e)[:200]}")
    draft = out["draft"]
    fields = studio.draft_to_fields(draft if isinstance(draft, dict) else {})
    if not fields:
        raise HTTPException(422, "未能从描述中起草出字段，请补充更具体的字段说明")
    return {"doc_type": (draft or {}).get("doc_type", ""),
            "fields": [f.model_dump() for f in fields],
            "provider_used": out["provider_used"]}


class EnrichBody(BaseModel):
    fields: list[dict]
    doc_type: str = ""
    provider: str | None = None


@router.post("/draft-enrich")
async def draft_enrich(body: EnrichBody):
    """LLM instruction enrichment (§5.1): expand terse field notes into full
    keyword->cleaning->format rules. Returns per-field instruction patches the
    editor merges as a draft — user reviews before saving. Burns tokens."""
    if not body.fields:
        raise HTTPException(400, "没有可补全的字段")
    await _warm_byok()
    try:
        out = await asyncio.to_thread(studio.enrich_fields, body.fields,
                                      body.doc_type, body.provider)
    except ProviderError as e:
        raise HTTPException(502, f"模型通道暂时不可用，请稍后重试：{str(e)[:200]}")
    draft = out["draft"] if isinstance(out["draft"], dict) else {}
    patches = []
    for f in draft.get("fields", []):
        if not isinstance(f, dict) or not f.get("name"):
            continue
        patches.append({"name": str(f["name"]),
                        "instruction": str(f.get("instruction") or ""),
                        "columns": [{"name": str(c["name"]),
                                     "instruction": str(c.get("instruction") or "")}
                                    for c in (f.get("columns") or [])
                                    if isinstance(c, dict) and c.get("name")]})
    if not patches:
        raise HTTPException(422, "模型未返回可用的说明补全")
    return {"fields": patches, "provider_used": out["provider_used"]}


@router.post("/draft-from-table")
async def draft_from_table(file: UploadFile = File(...)):
    """Rule import channel 3 (§5.1): spreadsheet/CSV field inventory -> drafted
    fields (pure parsing, zero tokens). Loose Chinese/English headers;
    a 所属明细表 column nests rows as table columns."""
    import csv
    import io

    name = (file.filename or "").lower()
    blob = await file.read()
    try:
        if name.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            ws = wb.active
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            wb.close()
        elif name.endswith((".csv", ".tsv", ".txt")):
            try:
                text = blob.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = blob.decode("gbk", errors="replace")
            delim = "\t" if name.endswith(".tsv") else ","
            rows = list(csv.reader(io.StringIO(text), delimiter=delim))
        else:
            raise HTTPException(400, "仅支持 .xlsx / .csv / .tsv 字段清单")
        fields = studio.fields_from_table(rows)
    except HTTPException:
        raise
    except ValueError as e:
        raise HTTPException(400, str(e))
    except Exception as e:
        raise HTTPException(400, f"表格解析失败：{e}")
    if not fields:
        raise HTTPException(422, "表格中没有可用的字段行")
    return {"fields": [f.model_dump() for f in fields]}


class DryRunBody(BaseModel):
    package: SkillPackage
    providers: list[str] = []


@router.post("/dry-run")
async def dry_run(file: UploadFile = File(...), package: str = Form(...),
                  providers: str = Form(default="")):
    """Try a package on one sample without creating a task; multiple providers
    run side-by-side (output/usage comparison — Unstract-validated UX)."""
    try:
        pkg = SkillPackage.model_validate_json(package)
    except Exception as e:
        raise HTTPException(400, f"invalid package json: {e}") from e
    plist = [p.strip() for p in providers.split(",") if p.strip()]
    sample_path, udr = await _parse_upload_kept(file, current_tenant(), pkg.parser)
    await _warm_byok()
    runs = await asyncio.to_thread(studio.dry_run, udr, pkg, plist or None,
                                   None, sample_path)
    return {"runs": runs}


@router.get("/{skill_code}/export", response_class=PlainTextResponse)
async def export_yaml(skill_code: str, version: int | None = None):
    """Skill package as YAML — cross-environment migration (§5.1)."""
    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        q = select(SkillVersion).where(SkillVersion.skill_code == skill_code,
                                       SkillVersion.tenant_id == tenant)
        q = q.where(SkillVersion.version == version) if version else \
            q.order_by(SkillVersion.version.desc())
        row = (await s.execute(q)).scalars().first()
        if row is None:
            raise HTTPException(404, "skill/version not found")
        return studio.package_to_yaml(SkillPackage(**row.package))


@router.post("/import", status_code=201)
async def import_yaml(file: UploadFile = File(...)):
    """YAML -> new skill (or new draft version if the code exists)."""
    tenant = current_tenant()
    try:
        pkg = studio.package_from_yaml((await file.read()).decode("utf-8"))
    except Exception as e:
        raise HTTPException(400, f"invalid package yaml: {e}") from e
    sf = session_factory()
    async with sf() as s:
        skill = await s.get(Skill, pkg.skill_code)
        if skill is not None and skill.tenant_id != tenant:
            raise HTTPException(409, "skill code taken by another tenant")
        if skill is None:
            await _enforce_skill_seat(s, tenant)   # new code = a new seat
            s.add(Skill(code=pkg.skill_code, tenant_id=tenant,
                        name=pkg.name or pkg.skill_code, kind=pkg.kind))
            next_ver = 1
        else:
            latest = (await s.execute(
                select(SkillVersion).where(SkillVersion.skill_code == pkg.skill_code)
                .order_by(SkillVersion.version.desc()))).scalars().first()
            next_ver = (latest.version + 1) if latest else 1
        s.add(SkillVersion(tenant_id=tenant, skill_code=pkg.skill_code,
                           version=next_ver, status="draft",
                           package=pkg.model_dump(), changelog="imported from YAML"))
        await s.commit()
    return {"skill_code": pkg.skill_code, "version": next_ver, "status": "draft"}


class GoldenAdd(BaseModel):
    expected: dict


@router.post("/{skill_code}/golden", status_code=201)
async def add_golden(skill_code: str, file: UploadFile = File(...),
                     expected: str = Form(...)):
    """Attach a golden sample (file + expected plain values) to a skill."""
    import json as _json
    tenant = current_tenant()
    blob = await file.read()
    digest = hashlib.sha256(blob).hexdigest()[:16]
    suffix = Path(file.filename or "golden").suffix.lower()
    store = Path(get_settings().data_dir) / "golden" / tenant / skill_code
    store.mkdir(parents=True, exist_ok=True)
    path = store / f"{digest}{suffix}"
    path.write_bytes(blob)
    sf = session_factory()
    async with sf() as s:
        if (await s.get(Skill, skill_code)) is None:
            raise HTTPException(404, "skill not found")
        s.add(GoldenSample(tenant_id=tenant, skill_code=skill_code,
                           storage_path=str(path),
                           expected=_json.loads(expected)))
        await s.commit()
    return {"skill_code": skill_code, "stored": path.name}


@router.post("/{skill_code}/versions/{version}/golden-check")
async def golden_check(skill_code: str, version: int):
    """Run this version over the skill's golden set and report field diffs —
    the publish gate evidence (PM item #4). Burns tokens: explicit call only."""
    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        row = (await s.execute(
            select(SkillVersion).where(SkillVersion.skill_code == skill_code,
                                       SkillVersion.version == version,
                                       SkillVersion.tenant_id == tenant))).scalar_one_or_none()
        if row is None:
            raise HTTPException(404, "version not found")
        goldens = (await s.execute(
            select(GoldenSample).where(GoldenSample.skill_code == skill_code,
                                       GoldenSample.tenant_id == tenant))).scalars().all()
    if not goldens:
        return {"samples": 0, "note": "no golden samples attached"}
    await _warm_byok()
    pkg = SkillPackage(**row.package)

    def _run():
        # the gate must exercise the parser the version actually pins, or it
        # certifies something production will not run
        pairs = [(parse_document(g.storage_path, pkg.parser), g.expected or {})
                 for g in goldens]
        return studio.golden_check(pkg, pairs)

    return await asyncio.to_thread(_run)


@router.post("/{skill_code}/versions/{version}/publish")
async def publish(skill_code: str, version: int):
    """Publish draft; previous published version is archived (rollback = republish it).
    TODO(M2): golden-set diff report gate before publish (PM item #4)."""
    tenant = current_tenant()
    sf = session_factory()
    async with sf() as s:
        target = (await s.execute(
            select(SkillVersion).where(SkillVersion.skill_code == skill_code,
                                       SkillVersion.version == version))).scalar_one_or_none()
        if target is None or target.tenant_id != tenant:
            raise HTTPException(404, "version not found")
        if target.status == "published":
            return {"skill_code": skill_code, "version": version, "status": "published"}
        current = (await s.execute(
            select(SkillVersion).where(SkillVersion.skill_code == skill_code,
                                       SkillVersion.status == "published"))).scalars().all()
        for c in current:
            c.status = "archived"
        target.status = "published"
        await s.commit()
    return {"skill_code": skill_code, "version": version, "status": "published"}


async def _warm_byok() -> None:
    """Token-burning studio routes run extraction in worker threads — the
    tenant BYOK keys AND the console-registered custom channels must both be
    warm before entering sync land (§11.10): resolution there cannot await."""
    from app.extraction import byok, custom_providers
    sf = session_factory()
    async with sf() as s:
        tenant = current_tenant()
        await byok.warm(s, tenant)
        await custom_providers.warm(s, tenant)
